import json
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from hunterbot.core.domain import Confidence, Finding, Severity, VulnerabilityCategory
from hunterbot.reporting import (
    DOCXReportGenerator,
    HTMLReportGenerator,
    JSONReportGenerator,
    MarkdownReportGenerator,
    PDFReportGenerator,
    XLSXReportGenerator,
    get_generator,
)
from hunterbot.reporting.ordering import sort_findings


def _finding(title: str, severity: Severity, **overrides) -> Finding:
    defaults = dict(
        title=title,
        category=VulnerabilityCategory.MISSING_SECURITY_HEADERS,
        severity=severity,
        confidence=Confidence.HIGH,
        description="A test finding description.",
        affected_asset="https://example.com",
        location="https://example.com/",
        scanner_name="test-scanner",
    )
    defaults.update(overrides)
    return Finding(**defaults)


_LOW = _finding("Low severity issue", Severity.LOW)
_CRITICAL = _finding(
    "Critical severity issue",
    Severity.CRITICAL,
    evidence="raw evidence",
    reproduction_steps="1. Do X\n2. Do Y",
    impact="Full compromise.",
    remediation="Patch it.",
    references=("https://example.com/advisory",),
    knowledge_source_id=7,
)


class TestSortFindings:
    def test_orders_most_severe_first(self) -> None:
        ordered = sort_findings([_LOW, _CRITICAL])
        assert [f.title for f in ordered] == ["Critical severity issue", "Low severity issue"]


class TestMarkdownReportGenerator:
    def test_generate_writes_file_with_all_sections(self, tmp_path: Path) -> None:
        output_path = tmp_path / "report.md"

        result = MarkdownReportGenerator().generate(findings=[_CRITICAL, _LOW], output_path=output_path)

        assert result == output_path
        content = output_path.read_text(encoding="utf-8")
        assert "# HunterBot Vulnerability Report" in content
        assert "Total findings: 2" in content
        assert "[CRITICAL] Critical severity issue" in content
        assert "raw evidence" in content
        assert "Patch it." in content
        assert "https://example.com/advisory" in content
        # Critical should appear before Low in the rendered output.
        assert content.index("Critical severity issue") < content.index("Low severity issue")

    def test_generate_with_no_findings(self, tmp_path: Path) -> None:
        output_path = tmp_path / "empty.md"
        MarkdownReportGenerator().generate(findings=[], output_path=output_path)
        assert "No findings" in output_path.read_text(encoding="utf-8")

    def test_creates_parent_directories(self, tmp_path: Path) -> None:
        output_path = tmp_path / "nested" / "dir" / "report.md"
        MarkdownReportGenerator().generate(findings=[], output_path=output_path)
        assert output_path.exists()


class TestJSONReportGenerator:
    def test_generate_produces_valid_json_with_all_findings(self, tmp_path: Path) -> None:
        output_path = tmp_path / "report.json"

        JSONReportGenerator().generate(findings=[_CRITICAL, _LOW], output_path=output_path)

        document = json.loads(output_path.read_text(encoding="utf-8"))
        assert document["total_findings"] == 2
        assert document["findings"][0]["title"] == "Critical severity issue"
        assert document["findings"][0]["references"] == ["https://example.com/advisory"]
        assert document["findings"][0]["knowledge_source_id"] == 7


class TestHTMLReportGenerator:
    def test_generate_escapes_and_includes_all_findings(self, tmp_path: Path) -> None:
        output_path = tmp_path / "report.html"
        malicious = _finding("<script>alert(1)</script>", Severity.HIGH)

        HTMLReportGenerator().generate(findings=[_CRITICAL, malicious], output_path=output_path)

        content = output_path.read_text(encoding="utf-8")
        assert "<!doctype html>" in content
        assert "<script>alert(1)</script>" not in content
        assert "&lt;script&gt;" in content
        assert "Critical severity issue" in content


class TestDOCXReportGenerator:
    def test_generate_includes_all_findings_and_fields(self, tmp_path: Path) -> None:
        output_path = tmp_path / "report.docx"

        DOCXReportGenerator().generate(findings=[_CRITICAL, _LOW], output_path=output_path)

        document = Document(str(output_path))
        full_text = "\n".join(p.text for p in document.paragraphs)
        table_text = "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        assert "HunterBot Vulnerability Report" in full_text
        assert "[CRITICAL] Critical severity issue" in full_text
        assert "[LOW] Low severity issue" in full_text
        assert "Patch it." in full_text
        assert "https://example.com/advisory" in full_text
        assert "https://example.com" in table_text  # affected asset, from the metadata table

    def test_generate_with_no_findings(self, tmp_path: Path) -> None:
        output_path = tmp_path / "empty.docx"
        DOCXReportGenerator().generate(findings=[], output_path=output_path)
        document = Document(str(output_path))
        assert any("No findings" in p.text for p in document.paragraphs)


class TestXLSXReportGenerator:
    def test_generate_writes_one_row_per_finding(self, tmp_path: Path) -> None:
        output_path = tmp_path / "report.xlsx"

        XLSXReportGenerator().generate(findings=[_CRITICAL, _LOW], output_path=output_path)

        workbook = load_workbook(str(output_path))
        sheet = workbook["Findings"]
        header = [cell.value for cell in sheet[1]]
        assert header[0] == "ID"
        assert header[1] == "Severity"
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 2
        # Most severe first.
        assert rows[0][1] == "critical"
        assert rows[0][2] == "Critical severity issue"
        assert rows[0][14] == "https://example.com/advisory"

    def test_generate_with_no_findings_writes_header_only(self, tmp_path: Path) -> None:
        output_path = tmp_path / "empty.xlsx"
        XLSXReportGenerator().generate(findings=[], output_path=output_path)
        workbook = load_workbook(str(output_path))
        sheet = workbook["Findings"]
        assert sheet.max_row == 1


class TestPDFReportGenerator:
    def test_generate_produces_readable_pdf_with_all_findings(self, tmp_path: Path) -> None:
        output_path = tmp_path / "report.pdf"

        PDFReportGenerator().generate(findings=[_CRITICAL, _LOW], output_path=output_path)

        reader = PdfReader(str(output_path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        assert "HunterBot Vulnerability Report" in text
        assert "Critical severity issue" in text
        assert "Low severity issue" in text
        assert "Patch it." in text
        assert "https://example.com/advisory" in text

    def test_generate_with_no_findings(self, tmp_path: Path) -> None:
        output_path = tmp_path / "empty.pdf"
        PDFReportGenerator().generate(findings=[], output_path=output_path)
        reader = PdfReader(str(output_path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        assert "No findings" in text

    def test_generate_handles_non_latin1_text_without_crashing(self, tmp_path: Path) -> None:
        output_path = tmp_path / "unicode.pdf"
        finding = _finding(
            "Unicode finding — café 中文",
            Severity.MEDIUM,
            description="Contains an em dash — and CJK 中文 characters.",
        )

        result = PDFReportGenerator().generate(findings=[finding], output_path=output_path)

        assert result == output_path
        assert output_path.exists()


class TestGetGenerator:
    def test_returns_matching_generator(self) -> None:
        assert isinstance(get_generator("markdown"), MarkdownReportGenerator)
        assert isinstance(get_generator("json"), JSONReportGenerator)
        assert isinstance(get_generator("html"), HTMLReportGenerator)
        assert isinstance(get_generator("pdf"), PDFReportGenerator)
        assert isinstance(get_generator("docx"), DOCXReportGenerator)
        assert isinstance(get_generator("xlsx"), XLSXReportGenerator)

    def test_unsupported_format_raises(self) -> None:
        with pytest.raises(ValueError, match="unsupported report format"):
            get_generator("csv")
