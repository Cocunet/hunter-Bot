from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from hunterbot.core.domain import Finding
from hunterbot.reporting.knowledge_labels import knowledge_source_label
from hunterbot.reporting.ordering import sort_findings

_METADATA_LABELS = (
    ("Category", lambda f: f.category.value),
    ("Confidence", lambda f: f.confidence.value),
    ("Affected asset", lambda f: f.affected_asset),
    ("Location", lambda f: f.location),
    ("Scanner", lambda f: f.scanner_name),
    ("Detected", lambda f: f.created_at.isoformat()),
)


def _add_finding(document: Document, finding: Finding, knowledge_titles: dict[int, str] | None) -> None:
    document.add_heading(f"[{finding.severity.value.upper()}] {finding.title}", level=2)

    table = document.add_table(rows=0, cols=2)
    table.style = "Light List"
    for label, getter in _METADATA_LABELS:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(getter(finding))
    label = knowledge_source_label(finding, knowledge_titles)
    if label is not None:
        row = table.add_row().cells
        row[0].text = "Knowledge source"
        row[1].text = label

    document.add_heading("Description", level=3)
    document.add_paragraph(finding.description)

    for label, value in (
        ("Evidence", finding.evidence),
        ("Reproduction", finding.reproduction_steps),
        ("Impact", finding.impact),
        ("Remediation", finding.remediation),
    ):
        if value:
            document.add_heading(label, level=3)
            document.add_paragraph(value)

    if finding.references:
        document.add_heading("References", level=3)
        for reference in finding.references:
            document.add_paragraph(reference, style="List Bullet")


class DOCXReportGenerator:
    format_name = "docx"

    def generate(
        self,
        *,
        findings: list[Finding],
        output_path: Path,
        knowledge_titles: dict[int, str] | None = None,
    ) -> Path:
        ordered = sort_findings(findings)
        document = Document()

        document.add_heading("HunterBot Vulnerability Report", level=1)
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
        generated_at = datetime.now(timezone.utc).isoformat()
        meta.add_run(f"Generated: {generated_at}    Total findings: {len(ordered)}").italic = True

        if not ordered:
            document.add_paragraph("No findings.")
        for finding in ordered:
            _add_finding(document, finding, knowledge_titles)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
        return output_path
